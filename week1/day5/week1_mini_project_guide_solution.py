from groq import Groq
from dotenv import load_dotenv
from pydantic import BaseModel
import time
import os
import json
from pathlib import Path

load_dotenv()

response_format={
    "type": "json_object"
}

my_api_key=os.getenv("GROQ_API_KEY")
using_model=os.getenv("GROQ_MODEL_USED")

client = Groq(api_key=my_api_key)

job_description="""
Description
Do you want to solve real customer problems through innovative technology? Do you enjoy working on scalable services in a collaborative team environment? Do you want to see your code directly impact millions of customers worldwide?

At Amazon, we hire the best minds in technology to innovate and build on behalf of our customers. Customer obsession is part of our company DNA, which has made us one of the world's most beloved brands.

Our Software Development Engineers (SDEs) use modern technology to solve complex problems while seeing their work's impact first-hand. The challenges SDEs solve at Amazon are meaningful and influence millions of customers, sellers, and products globally. We seek individuals passionate about creating new products, features, and services while managing ambiguity in an environment where development cycles are measured in weeks, not years.

At Amazon, we believe in ownership at every level. As an SDE-I, you'll own the entire lifecycle of your code - from design through deployment and ongoing operations. This ownership mindset, combined with our commitment to operational excellence, ensures we deliver the highest quality solutions for our customers.

We're looking for curious minds who think big and want to define tomorrow's technology. At Amazon, you'll grow into the high-impact engineer you know you can be, supported by a culture of learning and mentorship. Every day brings exciting new challenges and opportunities for personal growth.
Key job responsibilities
• Collaborate and communicate effectively with experienced cross-disciplinary Amazonians to design, build, and operate innovative products and services that delight our customers, while participating in technical discussions to drive solutions forward.
• Design and develop scalable solutions using cloud-native architectures and microservices in a large distributed computing environment.
• Participate in code reviews and contribute to technical documentation.
• Build and maintain resilient distributed systems that are scalable, fault-tolerant, and cost-effective.
• Leverage and contribute to the development of GenAI and AI-powered tools to enhance development productivity while staying current with emerging technologies.
• Write clean, maintainable code following best practices and design patterns.
• Work in an agile environment practicing CI/CD principles while participating in operational responsibilities including on-call duties.
• Demonstrate operational excellence through monitoring, troubleshooting, and resolving production issues.
Basic Qualifications
- Experience with at least one general-purpose programming language such as Java, Python, C++, C#, Go, Rust, or TypeScript
- Experience with data structure implementation, basic algorithm development, and/or object-oriented design principles
- Currently has, or is in the process of obtaining a bachelor’s degree in Computer Science, Computer Engineering, Data Science, Information Systems, or related STEM fields
- Must be 18 years of age of older
Preferred Qualifications
- Experience from previous technical internship(s) or demonstrated project experience
- Experience with one or more of the following: AI tools for development productivity, Cloud platforms (preferably AWS), Database systems (SQL and NoSQL), Contributing to open-source projects, Version control systems, Debugging and troubleshooting complex systems
- Demonstrated ability to learn and adapt to new technologies quickly
- Basic understanding of software development lifecycle (SDLC)
- Strong problem-solving and analytical skills
- Excellent written and verbal communication skills
"""

class JobDescription(BaseModel):
    role: str
    required_skills: list[str]
    preferred_skills: list[str]
    minimum_experience: float | None
    education_req: list[str]
    responsibilities: list[str]

job_desc_schema=JobDescription.model_json_schema()

system_prompt=f"""
You are an expert HR assistant.

Your job is to analyze job descriptions and extract
structured information from them.

Return ONLY valid JSON matching this schema:

{job_desc_schema}

IMPORTANT:
Do NOT return the schema itself.
Do NOT return fields like "properties", "title" or "type".
Fill the schema with actual information extracted from the job description.

If minimum experience is not mentioned, return null.
If information for a list is missing, return an empty list.
Do not invent information.
"""
system_message={
    "role": "system",
    "content": system_prompt
}

user_message={
    "role": "user",
    "content": job_description
}

response = client.chat.completions.create(
        model=using_model,
        messages=[system_message,user_message], 
        response_format=response_format
    )
json_op=response.choices[0].message.content

data=json.loads(json_op)
job_description=JobDescription(**data)

class Experience(BaseModel):
    company: str | None = None
    role: str| None = None
    duration: str| None = None
    skills: list[str] =[]
    description: str | None = None

class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None
    total_experience_years: float | None = None
    education: list[str] = []
    experience: list[Experience] = []
    projects: list[str] = []
    certifications: list[str] = []
    skills: list[str] = []

resume_schema = Resume.model_json_schema()

class MatchScore(BaseModel):
    score:float 
    final_verdict:dict

match_score_schema=MatchScore.model_json_schema()

from pypdf import PdfReader
from docx import Document

resume_folder = Path("resumes")
all_results=[]

def read_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        text += page.extract_text() + "\n"
    return text

def read_docx(file_path):
    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text

def read_resume(file_path):
    if file_path.suffix.lower() == ".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower() == ".docx":
        return read_docx(file_path)
    else:
        return None

def parsing_resume_text(resume_text):

    system_prompt=f"""
    I want you to format the text resume which i give you and give me JSOn object using attached schema only.
    {resume_schema}
    """    

    user_prompt=f"""
    you need yo format {resume_text} and give JSOn object based on attached schema.
    """

    system_message={
    "role":"system",
    "content":system_prompt
    }

    user_message={
        "role":"user",
        "content":user_prompt
    }

    response=client.chat.completions.create(
        model=using_model,
        messages=[system_message, user_message],
        response_format=response_format
    )

    json_op_of_resume=response.choices[0].message.content
    data_resume=json.loads(json_op_of_resume)
    resume_in_json_object=Resume(**data_resume)
    return resume_in_json_object

def MatchingResumes(jobDescription, candidateResume):

    system_prompt=f"""
    you are an expert resume matcher who matches resume and job description and 
    gives output as how the candiate is fit for the job and give percent match of it.
    give JSON object using {match_score_schema} schema.
    score variable of schema should contain percent match for the role fir of the candidate.
    final verdict should contain name, year of exp, skills, final say as to why candidate is a strong, medium, or weak fit for the role.
    """

    user_prompt=f"""
    The job descritpion is {jobDescription} and candidate resume is {candidateResume}
    """

    system_message ={
        "role":"system",
        "content":system_prompt
    }

    user_message = {
        "role" : "user",
        "content":user_prompt
    }

    messages=[system_message, user_message]

    response = client.chat.completions.create(
        model=using_model,
        messages=messages,
        response_format=response_format
    )

    matching_verdict = response.choices[0].message.content

    data_matching_verdict=json.loads(matching_verdict)
    matching_verdict_json=MatchScore(**data_matching_verdict)
    return matching_verdict_json

for file_path in resume_folder.iterdir():

    if file_path.suffix.lower() not in [".pdf", ".docx"]:
        continue
    
    print("\nProcessing:", file_path.name)
    resume_text = read_resume(file_path)
    parsed_resume_in_js_format=parsing_resume_text(resume_text)
    time.sleep(5)
    matching_json=MatchingResumes(job_description,parsed_resume_in_js_format)
    time.sleep(5)
    print(matching_json)